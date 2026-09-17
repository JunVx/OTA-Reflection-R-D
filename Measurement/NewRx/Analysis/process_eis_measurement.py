from __future__ import annotations

import math
import os
import shutil
import subprocess
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd


ANALYSIS_DIR = Path(__file__).resolve().parent
NEWRX_DIR = ANALYSIS_DIR.parent
FIG_DIR = ANALYSIS_DIR / "figures"
TABLE_DIR = ANALYSIS_DIR / "tables"
TEMP_COPY = Path(r"C:\Temp\newrx_tis_data.xlsx")


def fs_path(path: Path | str) -> str:
    p = Path(path)
    if not p.is_absolute():
        p = p.resolve()
    s = str(p)
    if os.name == "nt" and not s.startswith("\\\\?\\"):
        if s.startswith("\\\\"):
            return "\\\\?\\UNC\\" + s[2:]
        return "\\\\?\\" + s
    return s


def save_figure(fig: plt.Figure, path: Path, **kwargs) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fmt = path.suffix.lstrip(".").lower() or None
    with open(fs_path(path), "wb") as handle:
        fig.savefig(handle, format=fmt, **kwargs)


def locate_workbook() -> Path:
    files = [p for p in NEWRX_DIR.glob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        raise FileNotFoundError(f"No .xlsx file found in {NEWRX_DIR}")
    return files[0]


def prepare_short_workbook_copy() -> Path:
    source = locate_workbook()
    TEMP_COPY.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(fs_path(source), "rb") as src, open(TEMP_COPY, "wb") as dst:
            shutil.copyfileobj(src, dst)
    except PermissionError:
        command = (
            f"Copy-Item -LiteralPath '{source}' -Destination '{TEMP_COPY}' -Force"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", command], check=True)
    return TEMP_COPY


def dbm_to_mw(values_dbm: np.ndarray) -> np.ndarray:
    return 10.0 ** (np.asarray(values_dbm, dtype=float) / 10.0)


def mw_to_dbm(values_mw: np.ndarray) -> np.ndarray:
    return 10.0 * np.log10(np.maximum(np.asarray(values_mw, dtype=float), 1e-300))


def linear_average_dbm(values_dbm: np.ndarray, axis: int = 0) -> np.ndarray:
    return mw_to_dbm(np.mean(dbm_to_mw(values_dbm), axis=axis))


def setup_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 10,
            "axes.labelsize": 9,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "legend.fontsize": 8,
            "axes.linewidth": 0.8,
            "lines.linewidth": 1.5,
            "grid.linewidth": 0.5,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def read_total_eis() -> dict[int, pd.DataFrame]:
    workbook_path = prepare_short_workbook_copy()
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheet_name = next(name for name in workbook.sheetnames if "N78" in name and "TIS" in name)
    sheet = workbook[sheet_name]

    # Total EIS block: rows 49-61, grouped as 100M, 50M, 20M.
    groups = {
        100: {"angle": 2, "baseline": 3, "reflected": [4, 5, 6, 7, 8], "mean": 9},
        50: {"angle": 11, "baseline": 12, "reflected": [13, 14, 15, 16, 17], "mean": 18},
        20: {"angle": 20, "baseline": 21, "reflected": [22, 23, 24, 25, 26], "mean": 27},
    }
    reflected_labels = [
        "Center with reflector",
        "Front 10 cm",
        "Back 10 cm",
        "Left 10 cm",
        "Right 10 cm",
    ]

    result: dict[int, pd.DataFrame] = {}
    for bandwidth_mhz, columns in groups.items():
        rows = []
        for row_idx in range(49, 62):
            angle = sheet.cell(row_idx, columns["angle"]).value
            if not isinstance(angle, (int, float)):
                continue
            baseline = float(sheet.cell(row_idx, columns["baseline"]).value)
            reflected = [float(sheet.cell(row_idx, col).value) for col in columns["reflected"]]
            avg = float(sheet.cell(row_idx, columns["mean"]).value)
            row = {
                "bandwidth_mhz": bandwidth_mhz,
                "angle_deg": float(angle),
                "no_reflector_dbm": baseline,
                "spatial_average_dbm": avg,
            }
            row.update({label: value for label, value in zip(reflected_labels, reflected)})
            result[bandwidth_mhz] = result.get(bandwidth_mhz, [])
            rows.append(row)

        diff_rows = {}
        for row_idx in range(64, 77):
            angle = sheet.cell(row_idx, columns["angle"]).value
            if not isinstance(angle, (int, float)):
                continue
            before = [float(sheet.cell(row_idx, col).value) for col in columns["reflected"]]
            diff_rows[float(angle)] = {
                **{f"diff_before_{idx + 1}_db": value for idx, value in enumerate(before)},
                "diff_spatial_average_db": float(sheet.cell(row_idx, columns["mean"]).value),
            }

        df = pd.DataFrame(rows)
        for diff_col in [f"diff_before_{idx}_db" for idx in range(1, 6)] + ["diff_spatial_average_db"]:
            df[diff_col] = df["angle_deg"].map(lambda angle, col=diff_col: diff_rows.get(float(angle), {}).get(col, np.nan))
        result[bandwidth_mhz] = df
    return result


def closed_curve(df: pd.DataFrame, column: str) -> tuple[np.ndarray, np.ndarray]:
    plot_df = df.sort_values("angle_deg")
    theta = np.deg2rad(plot_df["angle_deg"].to_numpy(dtype=float))
    values = plot_df[column].to_numpy(dtype=float)
    return theta, values


def polar_radius(values_dbm: np.ndarray, radial_min_dbm: float) -> np.ndarray:
    return np.asarray(values_dbm, dtype=float) - radial_min_dbm


def plot_polar_single(df: pd.DataFrame, bandwidth_mhz: int, output_stem: str) -> None:
    reflected_cols = [
        "Center with reflector",
        "Front 10 cm",
        "Back 10 cm",
        "Left 10 cm",
        "Right 10 cm",
    ]
    all_values = [df["no_reflector_dbm"].to_numpy(), df["spatial_average_dbm"].to_numpy()]
    all_values.extend(df[col].to_numpy() for col in reflected_cols)
    all_values_flat = np.concatenate(all_values)
    radial_min = 5.0 * math.floor((float(np.nanmin(all_values_flat)) - 2.0) / 5.0)
    radial_max = 5.0 * math.ceil((float(np.nanmax(all_values_flat)) + 2.0) / 5.0)
    radial_ticks = np.arange(radial_min, radial_max + 0.1, 5.0)

    fig = plt.figure(figsize=(4.8, 4.9))
    ax = fig.add_subplot(111, projection="polar")
    ax.set_theta_zero_location("N")
    ax.set_theta_direction(-1)
    ax.set_title(f"N78 Total EIS, {bandwidth_mhz} MHz", pad=18)

    for idx, col in enumerate(reflected_cols):
        theta, values = closed_curve(df, col)
        ax.plot(
            theta,
            polar_radius(values, radial_min),
            color="0.28",
            lw=1.0,
            alpha=0.40,
            label="With reflector (5 positions)" if idx == 0 else None,
        )

    theta, baseline = closed_curve(df, "no_reflector_dbm")
    ax.plot(theta, polar_radius(baseline, radial_min), color="#2a9d55", lw=2.4, label="No reflector")

    theta, avg = closed_curve(df, "spatial_average_dbm")
    ax.plot(theta, polar_radius(avg, radial_min), color="#d62728", lw=2.2, label="Spatial average")

    ax.set_rlim(0, radial_max - radial_min)
    ax.set_rticks(radial_ticks - radial_min)
    ax.set_yticklabels([f"{tick:.0f}" for tick in radial_ticks])
    ax.set_rlabel_position(135)
    ax.grid(True, alpha=0.35)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.14), ncol=1, frameon=False)

    fig.subplots_adjust(bottom=0.24, top=0.88)
    save_figure(fig, FIG_DIR / f"{output_stem}.png", dpi=300, bbox_inches="tight")
    save_figure(fig, FIG_DIR / f"{output_stem}.pdf", bbox_inches="tight")
    plt.close(fig)


def plot_polar_combined(data: dict[int, pd.DataFrame]) -> None:
    bandwidths = [20, 50, 100]
    fig, axes = plt.subplots(1, 3, figsize=(10.4, 3.8), subplot_kw={"projection": "polar"})
    for ax, bandwidth in zip(axes, bandwidths):
        df = data[bandwidth]
        reflected_cols = [
            "Center with reflector",
            "Front 10 cm",
            "Back 10 cm",
            "Left 10 cm",
            "Right 10 cm",
        ]
        all_values = [df["no_reflector_dbm"].to_numpy(), df["spatial_average_dbm"].to_numpy()]
        all_values.extend(df[col].to_numpy() for col in reflected_cols)
        all_values_flat = np.concatenate(all_values)
        radial_min = 5.0 * math.floor((float(np.nanmin(all_values_flat)) - 2.0) / 5.0)
        radial_max = 5.0 * math.ceil((float(np.nanmax(all_values_flat)) + 2.0) / 5.0)
        radial_ticks = np.arange(radial_min, radial_max + 0.1, 5.0)
        ax.set_theta_zero_location("N")
        ax.set_theta_direction(-1)
        ax.set_title(f"{bandwidth} MHz", pad=14)
        for idx, col in enumerate(reflected_cols):
            theta, values = closed_curve(df, col)
            ax.plot(
                theta,
                polar_radius(values, radial_min),
                color="0.28",
                lw=0.9,
                alpha=0.35,
                label="With reflector (5 positions)" if idx == 0 else None,
            )
        theta, baseline = closed_curve(df, "no_reflector_dbm")
        ax.plot(theta, polar_radius(baseline, radial_min), color="#2a9d55", lw=2.2, label="No reflector")
        theta, avg = closed_curve(df, "spatial_average_dbm")
        ax.plot(theta, polar_radius(avg, radial_min), color="#d62728", lw=2.0, label="Spatial average")
        ax.set_rlim(0, radial_max - radial_min)
        ax.set_rticks(radial_ticks - radial_min)
        ax.set_yticklabels([f"{tick:.0f}" for tick in radial_ticks])
        ax.set_rlabel_position(135)
        ax.grid(True, alpha=0.35)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=3, frameon=False)
    fig.suptitle("N78 Total EIS Polar Measurement", y=1.02)
    fig.tight_layout(rect=(0, 0.10, 1, 0.98))
    save_figure(fig, FIG_DIR / "N78_Total_EIS_polar_all_bandwidths.png", dpi=300, bbox_inches="tight")
    save_figure(fig, FIG_DIR / "N78_Total_EIS_polar_all_bandwidths.pdf", bbox_inches="tight")
    plt.close(fig)


def build_summary(data: dict[int, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for bandwidth, df in sorted(data.items()):
        eval_df = df[df["angle_deg"] < 360].copy()
        diff_before_cols = [f"diff_before_{idx}_db" for idx in range(1, 6)]
        position_errors = eval_df[diff_before_cols].to_numpy(dtype=float)
        avg_error = eval_df["diff_spatial_average_db"].to_numpy(dtype=float)
        rows.append(
            {
                "Bandwidth (MHz)": bandwidth,
                "All reflected positions mean |error| (dB)": np.mean(np.abs(position_errors)),
                "All reflected positions max |error| (dB)": np.max(np.abs(position_errors)),
                "Spatial average mean |error| (dB)": np.mean(np.abs(avg_error)),
                "Spatial average max |error| (dB)": np.max(np.abs(avg_error)),
            }
        )
    return pd.DataFrame(rows)


def plot_error_summary(summary: pd.DataFrame) -> None:
    x = np.arange(len(summary))
    width = 0.36
    labels = [f"{int(value)}M" for value in summary["Bandwidth (MHz)"]]
    fig, ax = plt.subplots(figsize=(5.4, 3.0))
    ax.bar(
        x - width / 2,
        summary["All reflected positions mean |error| (dB)"],
        width,
        color="0.35",
        label="With reflector",
    )
    ax.bar(
        x + width / 2,
        summary["Spatial average mean |error| (dB)"],
        width,
        color="#d62728",
        label="Spatial average",
    )
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Mean |EIS error| (dB)")
    ax.set_title("N78 Total EIS Error Reduction")
    ax.grid(True, axis="y", alpha=0.28)
    ax.legend(frameon=False)
    fig.tight_layout()
    save_figure(fig, FIG_DIR / "N78_Total_EIS_error_summary.png", dpi=300, bbox_inches="tight")
    save_figure(fig, FIG_DIR / "N78_Total_EIS_error_summary.pdf", bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    setup_plot_style()
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    data = read_total_eis()
    all_rows = []
    for bandwidth, df in sorted(data.items()):
        df.to_csv(TABLE_DIR / f"N78_Total_EIS_{bandwidth}M_angles.csv", index=False, encoding="utf-8-sig")
        all_rows.append(df)
        plot_polar_single(df, bandwidth, f"N78_Total_EIS_{bandwidth}M_polar")
    all_data = pd.concat(all_rows, ignore_index=True)
    all_data.to_csv(TABLE_DIR / "N78_Total_EIS_angles_all_bandwidths.csv", index=False, encoding="utf-8-sig")
    plot_polar_combined(data)
    summary = build_summary(data)
    summary.to_csv(TABLE_DIR / "N78_Total_EIS_summary.csv", index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(TABLE_DIR / "N78_Total_EIS_analysis.xlsx", engine="openpyxl") as writer:
        all_data.to_excel(writer, sheet_name="Angle_Data", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
    plot_error_summary(summary)
    print(summary.round(3).to_string(index=False))
    print(f"Figures: {FIG_DIR}")
    print(f"Tables: {TABLE_DIR}")


if __name__ == "__main__":
    main()
